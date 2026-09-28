"""Fees page payment-date ("من تاريخ / إلى تاريخ") filtering.

The range selects money ACTUALLY COLLECTED in the period: the payment
allocation rows (``Revenue``) dated inside it, linked to their installment by
receipt number. Covers the screen, the Excel export and the print view, which
all share one filter implementation.
"""
import re
import unittest
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest import mock
from uuid import uuid4

from sqlalchemy import event

from app import create_app
from app.blueprints import fees as fees_mod
from app.models import (
    db, AcademicYear, AuditLog, FeeInstallment, FeeRecord, FeeType, Revenue,
    RevenueCategory, Role, School, Student, User,
)

OPTS = {'bypass_tenant_scope': True, 'include_all_years': True}
PASSWORD = 'Password123!'

SEP = {'payment_date_from': '2026-09-01', 'payment_date_to': '2026-09-30'}
OCT = {'payment_date_from': '2026-10-01', 'payment_date_to': '2026-10-31'}
SEP_OCT = {'payment_date_from': '2026-09-01', 'payment_date_to': '2026-10-31'}
AUG31_OCT = {'payment_date_from': '2026-08-31', 'payment_date_to': '2026-10-31'}
PAST_DUE = date(2026, 1, 10)


class FeesPaymentDateFilterTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.app = create_app('testing')
        cls.app.config.update(RATELIMIT_ENABLED=False, WTF_CSRF_ENABLED=False)
        with cls.app.app_context():
            cls.admin_role_id = Role.query.filter_by(name='school_admin').first().id

    # ── fixture ───────────────────────────────────────────────────────────────

    def setUp(self):
        self.sfx = uuid4().hex[:8]
        self.ids = {}
        with self.app.app_context():
            self._school('a')
            self._school('b')
            self._build_school_a()
            self._build_school_b()
            db.session.commit()

    def _school(self, key):
        s = self.sfx
        school = School(school_name=f'PDF {key} {s}', code=f'PD{key}{s}'[:20],
                        capacity=0, is_active=True)
        db.session.add(school)
        db.session.flush()
        year = AcademicYear(school_id=school.id, name=f'Y{key}{s}', is_current=True,
                            start_date=date(2026, 8, 1), end_date=date(2027, 7, 31))
        db.session.add(year)
        admin = User(username=f'pd{key}_{s}', email=f'pd{key}_{s}@t.test',
                     full_name=f'adm {key}', role_id=self.admin_role_id,
                     school_id=school.id, is_active=True)
        admin.set_password(PASSWORD)
        db.session.add(admin)
        cat = RevenueCategory(name='رسوم دراسية', school_id=school.id)
        db.session.add(cat)
        db.session.flush()
        self.ids.update({f'school_{key}': school.id, f'year_{key}': year.id,
                         f'admin_{key}': admin.id, f'admin_name_{key}': admin.username,
                         f'cat_{key}': cat.id})

    def _student(self, key, name, year_id=None, rfid=None):
        st = Student(student_id=f'ST{uuid4().hex[:10]}', full_name=f'{name} {self.sfx}',
                     school_id=self.ids[f'school_{key}'],
                     academic_year_id=year_id or self.ids[f'year_{key}'],
                     rfid_tag_id=rfid)
        db.session.add(st)
        db.session.flush()
        return st

    def _fee_type(self, key, name, year_id=None):
        ft = FeeType(name=f'{name} {self.sfx}', school_id=self.ids[f'school_{key}'],
                     academic_year_id=year_id or self.ids[f'year_{key}'])
        db.session.add(ft)
        db.session.flush()
        return ft

    def _fee(self, key, student, fee_type, amounts, year_id=None, due=PAST_DUE):
        year_id = year_id or self.ids[f'year_{key}']
        rec = FeeRecord(student_id=student.id, fee_type_id=fee_type.id,
                        academic_year_id=year_id, school_id=self.ids[f'school_{key}'],
                        total_amount=sum(amounts), discount=0)
        db.session.add(rec)
        db.session.flush()
        for no, amt in enumerate(amounts, 1):
            db.session.add(FeeInstallment(
                fee_record_id=rec.id, school_id=rec.school_id, academic_year_id=year_id,
                installment_no=no, amount=amt, due_date=due))
        db.session.flush()
        return rec

    def _insts(self, rec):
        return (FeeInstallment.query.execution_options(**OPTS)
                .filter_by(fee_record_id=rec.id)
                .order_by(FeeInstallment.installment_no).all())

    def _pay(self, key, rec, start_no, amount, paid):
        """Record a payment through the canonical recorder (same code as the
        Fees page and the Add Student wizard)."""
        cat = db.session.get(RevenueCategory, self.ids[f'cat_{key}'],
                             execution_options=OPTS)
        allocs = fees_mod.stage_installment_payment(
            self._insts(rec), start_no, Decimal(amount), fee_category=cat,
            student_name='x', paid_date=paid,
            collected_by=self.ids[f'admin_{key}'],
            rev_academic_year_id=rec.academic_year_id)
        db.session.flush()
        return allocs

    def _build_school_a(self):
        a = 'a'
        self.alpha = alpha = self._student(a, 'Alpha', rfid=f'CARD{self.sfx}')
        beta, gamma = self._student(a, 'Beta'), self._student(a, 'Gamma')
        delta, eps = self._student(a, 'Delta'), self._student(a, 'Eps')
        tuition, bus = self._fee_type(a, 'Tuition'), self._fee_type(a, 'Bus')
        self.ids.update(tuition=tuition.id, bus=bus.id, rfid=alpha.rfid_tag_id)

        # F1 — the accounting example: 100,000; 20,000 on Sep 10, 30,000 on Oct 5.
        f1 = self._fee(a, alpha, tuition, [50000, 50000])
        f1.installments.filter_by(installment_no=2).one().due_date = date(2027, 1, 10)
        self._pay(a, f1, 1, 20000, date(2026, 9, 10))
        self._pay(a, f1, 1, 30000, date(2026, 10, 5))
        # F2 — three September payments on ONE installment → fully paid.
        f2 = self._fee(a, beta, tuition, [30000])
        for d in (5, 20, 25):
            self._pay(a, f2, 1, 10000, date(2026, 9, d))
        # F3 — before-range (Aug 31) and after-range (Oct 1, installment 2) payments.
        f3 = self._fee(a, gamma, bus, [25000, 25000])
        self._pay(a, f3, 1, 25000, date(2026, 8, 31))
        self._pay(a, f3, 2, 10000, date(2026, 10, 1))
        # F4 — never paid.
        f4 = self._fee(a, alpha, bus, [40000])
        # F5 — LEGACY row: receipt number in the description, no [TXN:] tag.
        f5 = self._fee(a, delta, tuition, [20000])
        i5 = self._insts(f5)[0]
        i5.received_amount, i5.status = 7000, 'partial'
        i5.receipt_no = f'RCP-20260612-{uuid4().hex[:6].upper()}'
        db.session.add(Revenue(
            category_id=self.ids['cat_a'], school_id=f5.school_id,
            academic_year_id=f5.academic_year_id, amount=7000, date=date(2026, 9, 12),
            description=f'دفعة رسوم للطالب Delta - قسط #1 - {i5.receipt_no}'))
        # Unattributable pre-receipt-link payment row (no receipt number at all).
        db.session.add(Revenue(
            category_id=self.ids['cat_a'], school_id=f5.school_id,
            academic_year_id=f5.academic_year_id, amount=4000, date=date(2026, 9, 3),
            description='دفعة رسوم للطالب Delta - قسط #1'))
        # F6 — a September payment that was later refunded (inactive income).
        f6 = self._fee(a, eps, tuition, [9000])
        allocs = self._pay(a, f6, 1, 6000, date(2026, 9, 8))
        allocs[0]['revenue'].refunded_at = datetime(2026, 9, 9)
        # Other academic year of the SAME school, paid in September.
        old_year = AcademicYear(school_id=self.ids['school_a'], name=f'Old{self.sfx}',
                                is_current=False, start_date=date(2025, 8, 1),
                                end_date=date(2026, 7, 31))
        db.session.add(old_year)
        db.session.flush()
        old_type = self._fee_type(a, 'OldTuition', year_id=old_year.id)
        fo = self._fee(a, gamma, old_type, [11111], year_id=old_year.id)
        self._pay(a, fo, 1, 11111, date(2026, 9, 10))
        self.ids.update(f1=f1.id, f2=f2.id, f3=f3.id, f4=f4.id, f5=f5.id, f6=f6.id,
                        fo=fo.id, old_year=old_year.id,
                        f1_receipt=self._insts(f1)[0].receipt_no,
                        f3_inst2=self._insts(f3)[1].id, f4_inst1=self._insts(f4)[0].id,
                        f5_inst1=i5.id)

    def _build_school_b(self):
        b = 'b'
        fb = self._fee(b, self._student(b, 'Bravo'), self._fee_type(b, 'Tuition'), [99999])
        self._pay(b, fb, 1, 99999, date(2026, 9, 10))
        # A School B row that (wrongly) names School A's receipt number must never
        # be attributed to School A's installment.
        db.session.add(Revenue(
            category_id=self.ids['cat_b'], school_id=self.ids['school_b'],
            academic_year_id=self.ids['year_b'], amount=77777, date=date(2026, 9, 10),
            description=f"دفعة رسوم للطالب X - قسط #1 - {self.ids['f1_receipt']} [TXN:RCP-20260910-FFFFFF]"))
        self.ids['fb'] = fb.id

    def tearDown(self):
        with self.app.app_context():
            db.session.rollback()
            for key in ('a', 'b'):
                sid = self.ids[f'school_{key}']
                uids = [u.id for u in User.query.execution_options(**OPTS)
                        .filter_by(school_id=sid).all()]
                if uids:
                    AuditLog.query.execution_options(**OPTS).filter(
                        AuditLog.user_id.in_(uids)).delete(synchronize_session=False)
                for model in (Revenue, FeeInstallment, FeeRecord, FeeType, RevenueCategory,
                              Student, User, AcademicYear):
                    model.query.execution_options(**OPTS).filter_by(
                        school_id=sid).delete(synchronize_session=False)
                School.query.filter_by(id=sid).delete(synchronize_session=False)
            db.session.commit()
            db.session.remove()

    # ── helpers ───────────────────────────────────────────────────────────────

    def _client(self, key='a'):
        client = self.app.test_client()
        resp = client.post('/auth/login', data={'username': self.ids[f'admin_name_{key}'],
                                                'password': PASSWORD})
        self.assertEqual(resp.status_code, 302)
        return client

    def _ctx(self, path, key='a', **params):
        """GET a fees view and capture its template context."""
        captured = {}

        def fake_render(_tpl, **context):
            captured.update(context)
            return 'ok'

        client = self._client(key)
        with mock.patch('app.blueprints.fees.render_template', side_effect=fake_render):
            resp = client.get(path, query_string=params)
        self.assertEqual(resp.status_code, 200, resp.data[:500])
        return captured

    def _page(self, key='a', **params):
        ctx = self._ctx('/fees/', key, **params)
        if ctx['overdue_mode']:
            rows = {i.id for i in ctx['overdue_installments']}
        else:
            rows = {r.id for r, _ in ctx['fee_entries']}
        return ctx, rows

    def _period(self, ctx):
        return {k: int(v) for k, v in (ctx['period_paid'] or {}).items()}

    def _xlsx(self, key='a', **params):
        from openpyxl import load_workbook
        resp = self._client(key).get('/fees/export/excel', query_string=params)
        self.assertEqual(resp.status_code, 200)
        ws = load_workbook(BytesIO(resp.data)).active
        return [list(r) for r in ws.iter_rows(values_only=True)]

    # ── 1. no payment-date filter → existing behaviour ───────────────────────

    def test_01_no_date_filter_unchanged(self):
        ctx, rows = self._page()
        i = self.ids
        self.assertEqual(rows, {i['f1'], i['f2'], i['f3'], i['f4'], i['f5'], i['f6']})
        self.assertIsNone(ctx['period_paid'])
        self.assertIsNone(ctx['period_total'])
        self.assertFalse(ctx['period_active'])
        html = self._client().get('/fees/').get_data(as_text=True)
        self.assertIn('<th>المدفوع</th>', html)
        self.assertNotIn('المدفوع ضمن الفترة', html)
        self.assertNotIn('إجمالي المبالغ المستلمة ضمن الفترة', html)

    # ── 2–5. range semantics ─────────────────────────────────────────────────

    def test_02_from_only(self):
        ctx, rows = self._page(payment_date_from='2026-10-01')
        self.assertEqual(rows, {self.ids['f1'], self.ids['f3']})
        self.assertEqual(self._period(ctx), {self.ids['f1']: 30000, self.ids['f3']: 10000})

    def test_03_to_only(self):
        ctx, rows = self._page(payment_date_to='2026-08-31')
        self.assertEqual(rows, {self.ids['f3']})
        self.assertEqual(self._period(ctx), {self.ids['f3']: 25000})

    def test_04_both_inclusive(self):
        ctx, rows = self._page(**AUG31_OCT)
        i = self.ids
        self.assertEqual(rows, {i['f1'], i['f2'], i['f3'], i['f5']})
        self.assertEqual(self._period(ctx)[i['f3']], 35000)  # Aug 31 + Oct 1 both edges

    def test_05_same_day_range(self):
        ctx, rows = self._page(payment_date_from='2026-09-10', payment_date_to='2026-09-10')
        self.assertEqual(rows, {self.ids['f1']})
        self.assertEqual(int(ctx['period_total']), 20000)

    # ── 6–7. validation ───────────────────────────────────────────────────────

    def test_06_malformed_date_rejected_safely(self):
        for bad in ('abc', '2026-13-01', '2026/09/01', '20260901', '2026-02-30',
                    "2026-09-01' OR 1=1--"):
            ctx, rows = self._page(payment_date_from=bad)
            self.assertEqual(ctx['payment_date_error'], fees_mod.MSG_PAYMENT_DATE_INVALID, bad)
            self.assertEqual(rows, set(), bad)
            self.assertIsNone(ctx['period_total'])
        resp = self._client().get('/fees/', query_string={'payment_date_to': 'x'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn(fees_mod.MSG_PAYMENT_DATE_INVALID, resp.get_data(as_text=True))
        for path in ('/fees/export/excel', '/fees/print'):
            resp = self._client().get(path, query_string={'payment_date_from': 'x'})
            self.assertEqual(resp.status_code, 302, path)
            self.assertIn('/fees/?', resp.headers['Location'])

    def test_07_from_after_to_rejected(self):
        ctx, rows = self._page(payment_date_from='2026-09-30', payment_date_to='2026-09-01')
        self.assertEqual(ctx['payment_date_error'], fees_mod.MSG_PAYMENT_DATE_ORDER)
        self.assertEqual(rows, set())

    # ── 8–14, 28, 29. row / amount semantics ──────────────────────────────────

    def test_08_to_14_period_rows_and_amounts(self):
        ctx, rows = self._page(**SEP)
        i = self.ids
        # F4 (never paid), F3 (Aug 31 + Oct 1 only), F6 (refunded) excluded.
        self.assertEqual(rows, {i['f1'], i['f2'], i['f5']})
        self.assertEqual(self._period(ctx), {i['f1']: 20000, i['f2']: 30000, i['f5']: 7000})
        # F2 has three September payments but appears once.
        self.assertEqual([r.id for r, _ in ctx['fee_entries']].count(i['f2']), 1)
        self.assertEqual(int(ctx['period_total']), 57000)

    def test_15_accounting_example(self):
        f1 = self.ids['f1']
        self.assertEqual(self._period(self._page(**SEP)[0])[f1], 20000)
        self.assertEqual(self._period(self._page(**OCT)[0])[f1], 30000)
        ctx, _ = self._page(**SEP_OCT)
        self.assertEqual(self._period(ctx)[f1], 50000)
        self.assertEqual([r.id for r, _ in ctx['fee_entries']].count(f1), 1)
        self.assertIn(f1, self._page()[1])

    def test_28_partial_payment_period_amount(self):
        # F5 is a partially paid legacy installment (7,000 of 20,000).
        ctx, rows = self._page(**SEP)
        self.assertIn(self.ids['f5'], rows)
        self.assertEqual(self._period(ctx)[self.ids['f5']], 7000)

    # ── 15–19. combined with existing filters ────────────────────────────────

    def test_16_fee_type_plus_date(self):
        _, rows = self._page(fee_type=str(self.ids['tuition']), **SEP)
        self.assertEqual(rows, {self.ids['f1'], self.ids['f2'], self.ids['f5']})
        _, rows = self._page(fee_type=str(self.ids['bus']), **SEP)
        self.assertEqual(rows, set())
        ctx, rows = self._page(fee_type=str(self.ids['bus']), **AUG31_OCT)
        self.assertEqual(rows, {self.ids['f3']})
        self.assertEqual(int(ctx['period_total']), 35000)

    def test_17_status_plus_date(self):
        _, rows = self._page(payment_status='paid', **SEP)
        self.assertEqual(rows, {self.ids['f2']})
        _, rows = self._page(payment_status='unpaid', **SEP)
        self.assertEqual(rows, {self.ids['f1'], self.ids['f5']})

    def test_17b_overdue_plus_date(self):
        ctx, rows = self._page(payment_status='overdue', **SEP_OCT)
        i = self.ids
        # F4 inst 1 is overdue but received nothing in the period.
        self.assertEqual(rows, {i['f3_inst2'], i['f5_inst1']})
        self.assertEqual(self._period(ctx), {i['f3_inst2']: 10000, i['f5_inst1']: 7000})
        self.assertEqual(int(ctx['period_total']), 17000)
        _, rows = self._page(payment_status='overdue')
        self.assertIn(i['f4_inst1'], rows)

    def test_18_installment_plus_date(self):
        ctx, rows = self._page(installment='2', **AUG31_OCT)
        self.assertEqual(rows, {self.ids['f3']})
        self.assertEqual(self._period(ctx), {self.ids['f3']: 10000})
        _, rows = self._page(installment='1', **OCT)
        self.assertEqual(rows, {self.ids['f1']})

    def test_19_search_and_rfid_plus_date(self):
        _, rows = self._page(q='Alpha', **SEP)
        self.assertEqual(rows, {self.ids['f1']})
        _, rows = self._page(rfid=self.ids['rfid'], **SEP)
        self.assertEqual(rows, {self.ids['f1']})

    def test_20_all_filters_together(self):
        ctx, rows = self._page(q='Alpha', fee_type=str(self.ids['tuition']),
                               installment='1', payment_status='unpaid', **SEP)
        self.assertEqual(rows, {self.ids['f1']})
        self.assertEqual(int(ctx['period_total']), 20000)

    # ── isolation ─────────────────────────────────────────────────────────────

    def test_21_cross_school_excluded(self):
        ctx, rows = self._page(**SEP)
        self.assertNotIn(self.ids['fb'], rows)
        self.assertEqual(int(ctx['period_total']), 57000)   # neither 99,999 nor 77,777
        ctx_b, rows_b = self._page('b', **SEP)
        self.assertEqual(rows_b, {self.ids['fb']})
        self.assertEqual(int(ctx_b['period_total']), 99999)
        # The School B row naming School A's receipt links to no School B fee, so
        # it is reported as unattributable IN SCHOOL B — never credited to A.
        self.assertEqual(ctx_b['period_unlinked']['count'], 1)
        self.assertEqual(int(ctx_b['period_unlinked']['amount']), 77777)
        self.assertEqual(int(ctx['period_unlinked']['amount']), 4000)

    def test_22_academic_year_scope_unchanged(self):
        ctx, rows = self._page(**SEP)
        self.assertNotIn(self.ids['fo'], rows)
        self.assertNotIn(self.ids['fo'], ctx['period_paid'])
        _, rows_all = self._page()
        self.assertNotIn(self.ids['fo'], rows_all)

    def test_23_unattributable_rows_reported_not_counted(self):
        ctx, _ = self._page(**SEP)
        self.assertEqual(ctx['period_unlinked']['count'], 1)
        self.assertEqual(int(ctx['period_unlinked']['amount']), 4000)
        self.assertIsNone(self._page(**OCT)[0]['period_unlinked'])

    # ── UI / pagination / reset ──────────────────────────────────────────────

    def test_24_ui_labels_persist_and_reset(self):
        html = self._client().get('/fees/', query_string={**SEP, 'fee_type': str(self.ids['tuition'])}).get_data(as_text=True)
        self.assertIn('من تاريخ', html)
        self.assertIn('إلى تاريخ', html)
        self.assertIn('حسب تاريخ الدفع', html)
        self.assertIn('المدفوع الكلي', html)
        self.assertIn('المدفوع ضمن الفترة', html)
        self.assertIn('إجمالي المبالغ المستلمة ضمن الفترة', html)
        self.assertIn('57,000', html)
        self.assertRegex(html, r'name="payment_date_from"[^>]*value="2026-09-01"')
        self.assertRegex(html, r'name="payment_date_to"[^>]*value="2026-09-30"')
        # Export links carry the same filters.
        self.assertRegex(html, r'/fees/export/excel\?[^"]*payment_date_from=2026-09-01')
        self.assertRegex(html, r'/fees/print\?[^"]*payment_date_to=2026-09-30')
        # Reset → the bare page: no dates, no period.
        self.assertIn(f'href="/fees/"', html)
        ctx, _ = self._page()
        self.assertEqual((ctx['payment_date_from'], ctx['payment_date_to']), ('', ''))
        self.assertFalse(ctx['period_active'])

    def test_25_pagination_preserves_filters_and_no_n_plus_1(self):
        def revenue_statements(params):
            seen = []

            def before(conn, cursor, statement, *a):
                if 'revenues' in statement:
                    seen.append(statement)
            with self.app.app_context():
                engine = db.engine
            event.listen(engine, 'before_cursor_execute', before)
            try:
                resp = self._client().get('/fees/', query_string=params)
            finally:
                event.remove(engine, 'before_cursor_execute', before)
            self.assertEqual(resp.status_code, 200)
            return len(seen), resp.get_data(as_text=True)

        base_count, _ = revenue_statements(SEP)
        with self.app.app_context():
            tuition = db.session.get(FeeType, self.ids['tuition'], execution_options=OPTS)
            for n in range(21):
                st = self._student('a', f'Page{n:02d}')
                rec = self._fee('a', st, tuition, [1000])
                self._pay('a', rec, 1, 100, date(2026, 9, 15))
            db.session.commit()
        many_count, html = revenue_statements(SEP)
        # Constant number of payment queries regardless of row count (no N+1).
        self.assertEqual(base_count, many_count)
        # paginate COUNT + page SELECT (both embed the period subquery) + page
        # period map + period total + unattributable-rows check.
        self.assertLessEqual(many_count, 5)
        links = re.findall(r'href="(/fees/\?[^"]*page=2[^"]*)"', html)
        self.assertTrue(links)
        for href in links:
            self.assertIn('payment_date_from=2026-09-01', href)
            self.assertIn('payment_date_to=2026-09-30', href)
        ctx1, rows1 = self._page(**SEP)
        ctx2, rows2 = self._page(page=2, **SEP)
        self.assertEqual(ctx2['records'].total, 24)
        self.assertEqual(len(rows1) + len(rows2), 24)
        self.assertFalse(rows1 & rows2)
        self.assertEqual(int(ctx2['period_total']), 57000 + 2100)

    # ── exports ───────────────────────────────────────────────────────────────

    def test_26_excel_same_rows_and_period_amounts(self):
        page_ctx, page_rows = self._page(**SEP)
        sheet = self._xlsx(**SEP)
        header = sheet[0]
        self.assertIn('المدفوع الكلي', header)
        self.assertIn('المدفوع ضمن الفترة', header)
        col = header.index('المدفوع ضمن الفترة')
        data = [r for r in sheet[1:] if isinstance(r[0], int)]
        with self.app.app_context():
            names = {r.id: r.student.full_name for r in FeeRecord.query.execution_options(**OPTS)
                     .filter(FeeRecord.id.in_(page_rows)).all()}
        expected = {names[k]: v for k, v in self._period(page_ctx).items()}
        self.assertEqual({r[2]: int(r[col]) for r in data}, expected)
        summary = {r[1]: r[2] for r in sheet if r[1] == 'إجمالي المقبوض ضمن الفترة'}
        self.assertEqual(int(summary['إجمالي المقبوض ضمن الفترة']), 57000)
        # No date filter → the export keeps its existing columns.
        plain = self._xlsx()[0]
        self.assertIn('المدفوع', plain)
        self.assertNotIn('المدفوع ضمن الفترة', plain)

    def test_27_print_same_rows_and_total(self):
        page_ctx, page_rows = self._page(fee_type=str(self.ids['tuition']), **SEP)
        pctx = self._ctx('/fees/print', fee_type=str(self.ids['tuition']), **SEP)
        self.assertEqual({r.id for r, _ in pctx['fee_entries']}, page_rows)
        self.assertEqual(self._period(pctx), self._period(page_ctx))
        self.assertEqual(int(pctx['period_summary']['total']), int(page_ctx['period_total']))
        # Overdue print mirrors the overdue screen.
        octx, orows = self._page(payment_status='overdue', **SEP_OCT)
        pctx = self._ctx('/fees/print', payment_status='overdue', **SEP_OCT)
        self.assertEqual({i.id for i in pctx['overdue_installments']}, orows)
        self.assertEqual(int(pctx['period_summary']['total']), 17000)
        html = self._client().get('/fees/print', query_string=SEP).get_data(as_text=True)
        self.assertIn('المدفوع ضمن الفترة', html)
        self.assertIn('إجمالي المقبوض ضمن الفترة', html)
        self.assertIn('57,000', html)

    def test_30_export_total_equals_db_payment_sum(self):
        """Independent check: sum the qualifying Revenue rows straight from the
        DB and compare with every consumer's period total."""
        with self.app.app_context():
            rows = (Revenue.query.execution_options(**OPTS)
                    .filter(Revenue.school_id == self.ids['school_a'],
                            Revenue.academic_year_id == self.ids['year_a'],
                            Revenue.refunded_at.is_(None),
                            Revenue.date >= date(2026, 9, 1),
                            Revenue.date <= date(2026, 9, 30),
                            Revenue.description.like('%- RCP-%')).all())
            expected = int(sum(r.amount for r in rows))
        self.assertEqual(expected, 57000)
        self.assertEqual(int(self._page(**SEP)[0]['period_total']), expected)
        self.assertEqual(int(self._ctx('/fees/print', **SEP)['period_summary']['total']), expected)
        sheet = self._xlsx(**SEP)
        total = next(r[2] for r in sheet if r[1] == 'إجمالي المقبوض ضمن الفترة')
        self.assertEqual(int(total), expected)

    def test_31_pay_route_payment_is_counted_by_its_payment_date(self):
        """End-to-end through the real POST /fees/pay route."""
        client = self._client()
        with mock.patch('app.blueprints.fees.finalize_payment_notifications'):
            resp = client.post(f"/fees/pay/{self.ids['f4_inst1']}",
                               data={'received_amount': '1500', 'payment_method': 'cash',
                                     'paid_date': '2026-09-17'})
        self.assertEqual(resp.status_code, 200, resp.data[:300])
        self.assertEqual(resp.get_json()['status'], 'success')
        ctx, rows = self._page(**SEP)
        self.assertIn(self.ids['f4'], rows)
        self.assertEqual(self._period(ctx)[self.ids['f4']], 1500)
        self.assertEqual(int(ctx['period_total']), 57000 + 1500)
        self.assertNotIn(self.ids['f4'], self._page(**OCT)[1])


if __name__ == '__main__':
    unittest.main()
