"""Super Admin student permanent delete — two-case policy.

CASE A  no protected financial history  → physical delete of the student and
        every student-owned child row (incl. middle records + institute rows).
CASE B  paid / receipted installments or refund / cancellation events →
        the row is kept as an operational tombstone (status='deleted'); all
        financial rows are untouched; refused while a balance is still due.

Also covers authorization (Super Admin only, selected school required, other
school rejected), the tombstone being hidden from every school-facing list /
search / export / picker / by-ID page, the status allow-list, parent linking,
transaction rollback and the audit entries.
"""
import unittest
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from unittest import mock
from uuid import uuid4

from sqlalchemy import text

from app import create_app
from app.blueprints import fees as fees_mod
from app.models import (
    db, AcademicYear, AttendanceDevice, AuditLog, DeviceStudentMapping,
    FeeInstallment, FeeRecord, FeeReminderLog, FeeType, InstituteAttendanceRecord,
    InstituteAttendanceSession, InstituteGroupEnrollment, InstituteStudyGroup,
    RevenueCategory, Role, School, Student, StudentAttendance, StudentDocument,
    StudentMiddleRecord, StudentTransport, Subject, TransportRoute, User,
)
from app.utils.school_cleanup import cleanup_school_cascade

OPTS = {'bypass_tenant_scope': True, 'include_all_years': True}
PASSWORD = 'Password123!'

MSG_PHYSICAL = 'تم حذف الطالب نهائياً بنجاح.'
MSG_TOMBSTONE = 'تم حذف الطالب من النظام التشغيلي مع الاحتفاظ بسجلاته المالية السابقة.'
MSG_OUTSTANDING = 'لا يمكن حذف الطالب حالياً لوجود مبالغ مستحقة عليه'


class StudentPermanentDeleteTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.app = create_app('testing')
        cls.app.config.update(RATELIMIT_ENABLED=False, WTF_CSRF_ENABLED=False)
        with cls.app.app_context():
            roles = {r.name: r.id for r in Role.query.all()}
            cls.admin_role_id = roles['school_admin']
            cls.super_role_id = roles['super_admin']
            cls.parent_role_id = roles['parent']

    # ── fixture ───────────────────────────────────────────────────────────────

    def setUp(self):
        self.sfx = uuid4().hex[:8]
        self.ids = {}
        with self.app.app_context():
            self._school('a')
            self._school('b')
            sup = User(username=f'sdsup_{self.sfx}', email=f'sdsup_{self.sfx}@t.test',
                       full_name='super', role_id=self.super_role_id, school_id=None,
                       is_active=True)
            sup.set_password(PASSWORD)
            db.session.add(sup)
            parent = User(username=f'sdpar_{self.sfx}', email=f'sdpar_{self.sfx}@t.test',
                          full_name='parent', role_id=self.parent_role_id,
                          school_id=self.ids['school_a'], is_active=True)
            parent.set_password(PASSWORD)
            db.session.add(parent)
            db.session.flush()
            self.ids.update(super=sup.id, super_name=sup.username,
                            parent=parent.id, parent_name=parent.username)
            db.session.commit()

    def _school(self, key):
        s = self.sfx
        school = School(school_name=f'SD {key} {s}', code=f'SD{key}{s}'[:20],
                        capacity=0, is_active=True)
        db.session.add(school)
        db.session.flush()
        year = AcademicYear(school_id=school.id, name=f'Y{key}{s}', is_current=True,
                            start_date=date(2026, 8, 1), end_date=date(2027, 7, 31))
        db.session.add(year)
        admin = User(username=f'sd{key}_{s}', email=f'sd{key}_{s}@t.test',
                     full_name=f'adm {key}', role_id=self.admin_role_id,
                     school_id=school.id, is_active=True)
        admin.set_password(PASSWORD)
        db.session.add(admin)
        cat = RevenueCategory(name='رسوم دراسية', school_id=school.id)
        db.session.add(cat)
        db.session.flush()
        ft = FeeType(name=f'Tuition {s}', school_id=school.id, academic_year_id=year.id)
        db.session.add(ft)
        db.session.flush()
        self.ids.update({f'school_{key}': school.id, f'year_{key}': year.id,
                         f'admin_{key}': admin.id, f'admin_name_{key}': admin.username,
                         f'cat_{key}': cat.id, f'ft_{key}': ft.id})

    def tearDown(self):
        with self.app.app_context():
            db.session.rollback()
            for key in ('a', 'b'):
                db.session.execute(text("DELETE FROM student_middle_records"
                                        " WHERE school_id = :s"),
                                   {'s': self.ids[f'school_{key}']})
            AuditLog.query.execution_options(**OPTS).filter(
                AuditLog.user_id == self.ids['super']).delete(synchronize_session=False)
            for key in ('a', 'b'):
                cleanup_school_cascade(self.ids[f'school_{key}'])
            User.query.execution_options(**OPTS).filter_by(
                id=self.ids['super']).delete(synchronize_session=False)
            db.session.commit()
            db.session.remove()

    # ── data helpers (call inside an app context) ────────────────────────────

    def _student(self, key='a', name='Stu', status='active', **kw):
        st = Student(student_id=f'ST{uuid4().hex[:10]}', full_name=f'{name} {self.sfx}',
                     school_id=self.ids[f'school_{key}'],
                     academic_year_id=self.ids[f'year_{key}'], status=status, **kw)
        db.session.add(st)
        db.session.flush()
        return st

    def _fee(self, student, amounts, key='a'):
        rec = FeeRecord(student_id=student.id, fee_type_id=self.ids[f'ft_{key}'],
                        academic_year_id=self.ids[f'year_{key}'],
                        school_id=self.ids[f'school_{key}'],
                        total_amount=sum(amounts), discount=0)
        db.session.add(rec)
        db.session.flush()
        for no, amt in enumerate(amounts, 1):
            db.session.add(FeeInstallment(
                fee_record_id=rec.id, school_id=rec.school_id,
                academic_year_id=rec.academic_year_id, installment_no=no,
                amount=amt, due_date=date(2026, 9, 1)))
        db.session.flush()
        return rec

    def _insts(self, rec):
        return (FeeInstallment.query.execution_options(**OPTS)
                .filter_by(fee_record_id=rec.id)
                .order_by(FeeInstallment.installment_no).all())

    def _pay(self, rec, amount, key='a'):
        cat = db.session.get(RevenueCategory, self.ids[f'cat_{key}'],
                             execution_options=OPTS)
        fees_mod.stage_installment_payment(
            self._insts(rec), 1, Decimal(amount), fee_category=cat,
            student_name='x', paid_date=date(2026, 9, 10),
            collected_by=self.ids[f'admin_{key}'],
            rev_academic_year_id=rec.academic_year_id)
        db.session.flush()

    def _link_parent(self, student):
        db.session.execute(text("INSERT INTO parent_students (user_id, student_id)"
                                " VALUES (:u, :s)"),
                           {'u': self.ids['parent'], 's': student.id})

    def _operational_links(self, student, rec=None, key='a'):
        """parent link + transport + device mapping (+ reminder log if rec)."""
        sid_school = self.ids[f'school_{key}']
        self._link_parent(student)
        route = TransportRoute(school_id=sid_school, name=f'R{uuid4().hex[:6]}',
                               driver_name='d', driver_phone='1', vehicle_type='bus',
                               vehicle_number='1')
        dev = AttendanceDevice(school_id=sid_school, name=f'D{uuid4().hex[:6]}',
                               device_scope='students', ip_address='127.0.0.1',
                               password='x', device_sn=f'SN-{uuid4().hex[:10]}')
        db.session.add_all([route, dev])
        db.session.flush()
        db.session.add(StudentTransport(school_id=sid_school, route_id=route.id,
                                        student_id=student.id))
        db.session.add(DeviceStudentMapping(school_id=sid_school, device_id=dev.id,
                                            employee_no_string=str(student.id),
                                            student_id=student.id, is_active=True))
        if rec is not None:
            db.session.add(FeeReminderLog(
                school_id=sid_school, academic_year_id=rec.academic_year_id,
                student_id=student.id, installment_id=self._insts(rec)[0].id,
                parent_user_id=self.ids['parent'], due_date=date(2026, 9, 1)))
        db.session.flush()

    def _count(self, table, sid, col='student_id'):
        return db.session.execute(
            text(f"SELECT COUNT(*) FROM {table} WHERE {col} = :sid"), {'sid': sid}).scalar()

    def _fin_snapshot(self, key='a'):
        """Every financial figure of the school, straight from the tables."""
        s = {'s': self.ids[f'school_{key}']}
        q = lambda sql: [tuple(r) for r in db.session.execute(text(sql), s).fetchall()]
        return {
            'revenues': q("SELECT id, amount, date, description, refunded_at, refund_event_id"
                          " FROM revenues WHERE school_id = :s ORDER BY id"),
            'active_revenue_total': q("SELECT COALESCE(SUM(amount), 0) FROM revenues"
                                      " WHERE school_id = :s AND refunded_at IS NULL"),
            'fee_records': q("SELECT id, student_id, total_amount, discount, cancelled_at"
                             " FROM fee_records WHERE school_id = :s ORDER BY id"),
            'installments': q("SELECT id, fee_record_id, amount, received_amount, receipt_no,"
                              " status, paid_date FROM fee_installments"
                              " WHERE school_id = :s ORDER BY id"),
            'paid_total': q("SELECT COALESCE(SUM(fi.received_amount), 0) FROM fee_installments fi"
                            " JOIN fee_records fr ON fr.id = fi.fee_record_id"
                            " WHERE fr.school_id = :s AND fr.cancelled_at IS NULL"),
            'refund_events': q("SELECT id, student_id, fee_record_id, event_type, amount"
                               " FROM fee_refund_events WHERE school_id = :s ORDER BY id"),
        }

    # ── HTTP helpers ─────────────────────────────────────────────────────────

    def _login(self, username):
        client = self.app.test_client()
        resp = client.post('/auth/login', data={'username': username, 'password': PASSWORD})
        self.assertEqual(resp.status_code, 302)
        return client

    def _admin(self, key='a'):
        return self._login(self.ids[f'admin_name_{key}'])

    def _super(self, key='a'):
        client = self._login(self.ids['super_name'])
        if key:
            with client.session_transaction() as sess:
                sess['active_school_id'] = self.ids[f'school_{key}']
        return client

    def _flashes(self, client):
        with client.session_transaction() as sess:
            return [m for _cat, m in sess.get('_flashes', [])]

    def _delete(self, client, sid):
        return client.post(f'/students/{sid}/delete')

    def _student_row(self, sid):
        return db.session.execute(
            text("SELECT status, section_id, rfid_tag_id, building_id, residential_area_id"
                 " FROM students WHERE id = :sid"), {'sid': sid}).fetchone()

    def _make_tombstone(self):
        """Student with a fully paid fee → tombstoned through the real route."""
        with self.app.app_context():
            st = self._student(name='Tomb', rfid_tag_id=f'C{self.sfx}')
            rec = self._fee(st, [5000])
            self._pay(rec, 5000)
            self._link_parent(st)
            db.session.commit()
            sid = st.id
        client = self._super()
        self._delete(client, sid)
        with self.app.app_context():
            self.assertEqual(self._student_row(sid).status, 'deleted')
        return sid

    # ── 1–3. authorization + school scope ─────────────────────────────────────

    def test_01_school_admin_cannot_permanently_delete(self):
        with self.app.app_context():
            sid = self._student().id
            db.session.commit()
        client = self._admin()
        resp = self._delete(client, sid)
        self.assertIn(resp.status_code, (302, 403))
        with self.app.app_context():
            self.assertEqual(self._student_row(sid).status, 'active')

    def test_02_super_admin_without_selected_school_rejected(self):
        with self.app.app_context():
            sid = self._student().id
            db.session.commit()
        client = self._super(key=None)
        resp = self._delete(client, sid)
        self.assertEqual(resp.status_code, 302)
        self.assertIn('يرجى اختيار المدرسة أولاً', ' '.join(self._flashes(client)))
        with self.app.app_context():
            self.assertEqual(self._student_row(sid).status, 'active')

    def test_03_super_admin_cannot_delete_student_of_other_selected_school(self):
        with self.app.app_context():
            sid = self._student('a').id
            db.session.commit()
        client = self._super('b')
        resp = self._delete(client, sid)
        self.assertEqual(resp.status_code, 404)
        with self.app.app_context():
            self.assertEqual(self._student_row(sid).status, 'active')

    # ── 4. CASE A — physical delete ───────────────────────────────────────────

    def test_04_no_financial_history_physical_delete(self):
        with self.app.app_context():
            a, ya = self.ids['school_a'], self.ids['year_a']
            st = self._student(name='Clean')
            rec = self._fee(st, [3000, 3000])          # never paid, no receipt
            self._operational_links(st, rec)
            db.session.add(StudentAttendance(student_id=st.id, school_id=a,
                                             academic_year_id=ya, status='present'))
            db.session.add(StudentDocument(student_id=st.id, school_id=a,
                                           academic_year_id=ya, document_type='id',
                                           file_path='x.pdf'))
            db.session.add(StudentMiddleRecord(school_id=a, student_id=st.id,
                                               academic_year_id=ya,
                                               snap_full_name=st.full_name,
                                               subject_grades={}))
            subj = Subject(school_id=a, academic_year_id=ya, name='Math',
                           code=f'M{self.sfx}')
            db.session.add(subj)
            db.session.flush()
            grp = InstituteStudyGroup(school_id=a, academic_year_id=ya,
                                      subject_id=subj.id, name='G1', is_active=True)
            db.session.add(grp)
            db.session.flush()
            db.session.add(InstituteGroupEnrollment(school_id=a, group_id=grp.id,
                                                    student_id=st.id))
            sess = InstituteAttendanceSession(
                school_id=a, academic_year_id=ya, group_id=grp.id,
                session_date=date(2026, 9, 6), start_time=time(16, 0),
                end_time=time(17, 0), status='not_recorded')
            db.session.add(sess)
            db.session.flush()
            db.session.add(InstituteAttendanceRecord(
                school_id=a, session_id=sess.id, student_id=st.id, status='present',
                source='manual_admin', recorded_at=datetime.utcnow()))
            db.session.commit()
            sid, code = st.id, st.student_id
            before = self._fin_snapshot()

        client = self._super()
        resp = self._delete(client, sid)
        self.assertEqual(resp.status_code, 302)
        self.assertIn(MSG_PHYSICAL, self._flashes(client))

        with self.app.app_context():
            self.assertIsNone(self._student_row(sid))
            for tbl in ('parent_students', 'fee_records', 'fee_reminder_logs',
                        'student_attendance', 'student_documents', 'student_middle_records',
                        'institute_group_enrollments', 'institute_attendance_records',
                        'student_transport', 'device_student_mappings'):
                self.assertEqual(self._count(tbl, sid), 0, tbl)
            after = self._fin_snapshot()
            # Revenues untouched (none belonged to this student anyway).
            self.assertEqual(after['revenues'], before['revenues'])
            self.assertEqual(after['refund_events'], before['refund_events'])
            log = (AuditLog.query.execution_options(**OPTS)
                   .filter_by(resource='student', resource_id=sid,
                              user_id=self.ids['super']).one())
            self.assertIn('outcome=physical_delete', log.details)
            self.assertIn(f'student_code={code}', log.details)
            self.assertIn(f"school_id={self.ids['school_a']}", log.details)

    # ── 5. CASE B — paid installment → tombstone ──────────────────────────────

    def test_05_paid_installment_tombstones_and_preserves_finance(self):
        with self.app.app_context():
            st = self._student(name='Paid', rfid_tag_id=f'C{self.sfx}')
            rec = self._fee(st, [4000, 6000])
            self._pay(rec, 10000)                       # fully paid → no balance due
            self._operational_links(st, rec)
            db.session.add(StudentAttendance(student_id=st.id,
                                             school_id=self.ids['school_a'],
                                             academic_year_id=self.ids['year_a'],
                                             status='present'))
            db.session.commit()
            sid, code = st.id, st.student_id
            receipts = [i.receipt_no for i in self._insts(rec)]
            self.assertTrue(all(receipts))
            before = self._fin_snapshot()

        client = self._super()
        resp = self._delete(client, sid)
        self.assertEqual(resp.status_code, 302)
        self.assertIn(MSG_TOMBSTONE, self._flashes(client))

        with self.app.app_context():
            row = self._student_row(sid)
            self.assertEqual(tuple(row), ('deleted', None, None, None, None))
            name = db.session.execute(text("SELECT full_name, student_id FROM students"
                                           " WHERE id = :s"), {'s': sid}).fetchone()
            self.assertEqual(name.student_id, code)                # not anonymized
            # Financial history byte-identical (records, installments incl.
            # received_amount + receipt_no, revenues, totals).
            self.assertEqual(self._fin_snapshot(), before)
            # Operational links removed.
            for tbl in ('parent_students', 'fee_reminder_logs', 'student_transport',
                        'device_student_mappings'):
                self.assertEqual(self._count(tbl, sid), 0, tbl)
            # Academic history kept.
            self.assertEqual(self._count('student_attendance', sid), 1)
            log = (AuditLog.query.execution_options(**OPTS)
                   .filter_by(resource='student', resource_id=sid,
                              user_id=self.ids['super']).one())
            self.assertIn('outcome=financial_tombstone', log.details)
            self.assertIn(f'student_code={code}', log.details)

    # ── 6. CASE B — refund / cancellation history → tombstone ────────────────

    def test_06_refund_history_tombstones_and_preserves_event(self):
        with self.app.app_context():
            st = self._student(name='Refund')
            rec = self._fee(st, [3000])
            self._pay(rec, 3000)
            fees_mod.cancel_fee_record(rec, self._insts(rec), reason='test cancel')
            db.session.commit()
            sid = st.id
            before = self._fin_snapshot()
            self.assertEqual(len(before['refund_events']), 1)

        client = self._super()
        self._delete(client, sid)
        self.assertIn(MSG_TOMBSTONE, self._flashes(client))
        with self.app.app_context():
            self.assertEqual(self._student_row(sid).status, 'deleted')
            self.assertEqual(self._fin_snapshot(), before)

        # Refund history (financial view) still shows the event for the school.
        resp = self._admin().get('/fees/refunds-history')
        self.assertEqual(resp.status_code, 200)
        self.assertIn(f'Refund {self.sfx}', resp.get_data(as_text=True))

    # ── 7. outstanding balance blocks ─────────────────────────────────────────

    def test_07_outstanding_balance_blocks_and_changes_nothing(self):
        with self.app.app_context():
            st = self._student(name='Owes')
            rec = self._fee(st, [5000, 5000])
            self._pay(rec, 2000)                        # protected + balance due
            self._link_parent(st)
            db.session.commit()
            sid = st.id
            before = self._fin_snapshot()

        client = self._super()
        resp = self._delete(client, sid)
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(any(MSG_OUTSTANDING in m for m in self._flashes(client)))
        with self.app.app_context():
            self.assertEqual(self._student_row(sid).status, 'active')
            self.assertEqual(self._fin_snapshot(), before)
            self.assertEqual(self._count('parent_students', sid), 1)
            self.assertEqual(AuditLog.query.execution_options(**OPTS)
                             .filter_by(resource='student', resource_id=sid).count(), 0)

    # ── 8. hidden from lists / search / export / pickers / parent linking ────

    def test_08_tombstone_hidden_everywhere(self):
        tomb = self._make_tombstone()
        with self.app.app_context():
            active = self._student(name='Live').id
            archived = self._student(name='Arch', status='archived').id
            db.session.commit()
        admin = self._admin()

        def index_ids(**params):
            captured = {}
            with mock.patch('app.blueprints.students.render_template',
                            side_effect=lambda _t, **c: captured.update(c) or 'ok'):
                resp = admin.get('/students/', query_string=params)
            self.assertEqual(resp.status_code, 200)
            return {s.id for s in captured['students'].items}

        def search_ids(**params):
            resp = admin.get('/students/search', query_string={'q': self.sfx, **params})
            self.assertEqual(resp.status_code, 200)
            return {i['id'] for i in resp.get_json()['items']}

        for fetch in (index_ids, search_ids):
            default = fetch()
            self.assertIn(active, default)
            self.assertNotIn(tomb, default)
            self.assertNotIn(archived, default)
            arch = fetch(status='archived')
            self.assertEqual(arch & {tomb, active, archived}, {archived})
            self.assertEqual(fetch(status='deleted') & {tomb}, set())

        # Export — normal and crafted status.
        from openpyxl import load_workbook
        for status in ('active', 'archived', 'deleted'):
            resp = admin.get('/students/export/excel', query_string={'status': status})
            self.assertEqual(resp.status_code, 200)
            cells = [str(c) for r in load_workbook(BytesIO(resp.data)).active
                     .iter_rows(values_only=True) for c in r if c is not None]
            self.assertFalse(any(f'Tomb {self.sfx}' in c for c in cells), status)

        # Fee statement search + registration-record pickers.
        resp = admin.get('/fees/students/search-statement', query_string={'q': self.sfx})
        found = {r['id'] for r in resp.get_json()['results']}
        self.assertNotIn(tomb, found)
        self.assertIn(archived, found)
        resp = admin.get('/student-registration-records/search-students',
                         query_string={'q': self.sfx})
        self.assertNotIn(tomb, {r['id'] for r in resp.get_json()})
        self.assertEqual(admin.get(f'/student-registration-records/student-data/{tomb}')
                         .status_code, 404)

        # Parent linking — the parent's link was removed by the tombstone; a
        # crafted re-link request cannot bring it back.
        with self.app.app_context():
            self.assertEqual(self._count('parent_students', tomb), 0)
        resp = admin.post(f"/admin/users/{self.ids['parent']}/edit", data={
            'username': self.ids['parent_name'], 'full_name': 'parent', 'is_active': '1',
            'student_ids': [str(tomb), str(active)]})
        self.assertEqual(resp.status_code, 302)
        with self.app.app_context():
            linked = {r[0] for r in db.session.execute(
                text("SELECT student_id FROM parent_students WHERE user_id = :u"),
                {'u': self.ids['parent']})}
            self.assertEqual(linked, {active})

    # ── 9b. attendance profile ───────────────────────────────────────────────

    def test_09b_tombstone_attendance_profile_404_history_kept(self):
        with self.app.app_context():
            live = self._student(name='LiveAtt')
            st = self._student(name='TombAtt')
            for s in (live, st):
                db.session.add(StudentAttendance(student_id=s.id,
                                                 school_id=self.ids['school_a'],
                                                 academic_year_id=self.ids['year_a'],
                                                 date=date(2026, 9, 7), status='present'))
            rec = self._fee(st, [2000])
            self._pay(rec, 2000)
            db.session.commit()
            live_id, tomb = live.id, st.id
        self._delete(self._super(), tomb)
        with self.app.app_context():
            self.assertEqual(self._student_row(tomb).status, 'deleted')
            before = self._fin_snapshot()

        admin = self._admin()
        self.assertEqual(admin.get(f'/attendance/student/{live_id}/profile').status_code, 200)
        self.assertEqual(admin.get(f'/attendance/student/{tomb}/profile').status_code, 404)
        self.assertEqual(self._super().get(f'/attendance/student/{tomb}/profile').status_code,
                         404)
        # Attendance rows themselves are untouched; financial views still open.
        with self.app.app_context():
            self.assertEqual(self._count('student_attendance', tomb), 1)
            self.assertEqual(self._fin_snapshot(), before)
        self.assertEqual(admin.get(f'/fees/student/{tomb}/statement').status_code, 200)
        self.assertEqual(admin.get('/fees/refunds-history').status_code, 200)

    # ── 9–10. direct URL + no reactivation ───────────────────────────────────

    def test_09_10_tombstone_unreachable_by_url_and_not_restorable(self):
        tomb = self._make_tombstone()
        admin = self._admin()
        self.assertEqual(admin.get(f'/students/{tomb}').status_code, 404)
        self.assertEqual(admin.get(f'/students/{tomb}/edit').status_code, 404)
        self.assertEqual(admin.post(f'/students/{tomb}/edit',
                                    data={'full_name': 'X', 'status': 'active'}).status_code, 404)
        self.assertEqual(admin.post(f'/students/{tomb}/archive',
                                    data={'status': 'active'}).status_code, 404)
        self.assertEqual(admin.post(f'/students/{tomb}/unlink-parent',
                                    data={'parent_user_id': self.ids['parent']}).status_code, 404)
        self.assertEqual(admin.get(f'/students/{tomb}/create-success').status_code, 404)
        # Hidden in the school UI for the Super Admin too; a second delete is a 404.
        sup = self._super()
        self.assertEqual(sup.get(f'/students/{tomb}').status_code, 404)
        self.assertEqual(self._delete(sup, tomb).status_code, 404)
        with self.app.app_context():
            self.assertEqual(self._student_row(tomb).status, 'deleted')

    def test_10b_status_allow_list_on_edit_and_archive(self):
        with self.app.app_context():
            sid = self._student(name='Allow').id
            db.session.commit()
        admin = self._admin()
        resp = admin.post(f'/students/{sid}/edit',
                          data={'full_name': f'Renamed {self.sfx}', 'status': 'deleted'})
        self.assertEqual(resp.status_code, 302)
        with self.app.app_context():
            row = db.session.execute(text("SELECT full_name, status FROM students"
                                          " WHERE id = :s"), {'s': sid}).fetchone()
            self.assertEqual(row.full_name, f'Renamed {self.sfx}')   # save path reached
            self.assertEqual(row.status, 'active')                     # 'deleted' refused
        admin.post(f'/students/{sid}/archive', data={'status': 'deleted'})
        with self.app.app_context():
            self.assertEqual(self._student_row(sid).status, 'archived')

    # ── 11. historical totals unchanged ──────────────────────────────────────

    def test_11_financial_totals_and_statement_unchanged(self):
        with self.app.app_context():
            other = self._student(name='Other')
            orec = self._fee(other, [7000])
            self._pay(orec, 1500)
            st = self._student(name='Hist')
            rec = self._fee(st, [2500, 2500])
            self._pay(rec, 5000)
            db.session.commit()
            sid = st.id
            receipt = self._insts(rec)[0].receipt_no
            before = self._fin_snapshot()
        statement_before = self._admin().get(f'/fees/student/{sid}/statement')
        self.assertEqual(statement_before.status_code, 200)

        self._delete(self._super(), sid)

        with self.app.app_context():
            self.assertEqual(self._student_row(sid).status, 'deleted')
            after = self._fin_snapshot()
            self.assertEqual(after, before)
            self.assertEqual(after['paid_total'], [(Decimal('6500.00'),)])
        # Historical statement remains retrievable by its direct financial URL.
        statement = self._admin().get(f'/fees/student/{sid}/statement')
        self.assertEqual(statement.status_code, 200)
        self.assertIn(receipt, statement.get_data(as_text=True))

    # ── 15. transaction safety ───────────────────────────────────────────────

    def test_15_failure_rolls_back_everything(self):
        with self.app.app_context():
            st = self._student(name='Roll', rfid_tag_id=f'R{self.sfx}')
            rec = self._fee(st, [1000])
            self._pay(rec, 1000)
            self._link_parent(st)
            db.session.commit()
            sid = st.id
            before = self._fin_snapshot()
        client = self._super()
        with mock.patch('app.utils.student_removal._TOMBSTONE_UNLINK_TABLES',
                        ('fee_reminder_logs', 'no_such_table_xyz')):
            resp = self._delete(client, sid)
        self.assertEqual(resp.status_code, 302)
        self.assertIn('لم يتم حذف أي شيء', ' '.join(self._flashes(client)))
        with self.app.app_context():
            row = self._student_row(sid)
            self.assertEqual(row.status, 'active')
            self.assertEqual(row.rfid_tag_id, f'R{self.sfx}')
            self.assertEqual(self._count('parent_students', sid), 1)
            self.assertEqual(self._fin_snapshot(), before)


if __name__ == '__main__':
    unittest.main()
