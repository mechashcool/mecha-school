"""
Students Excel export = exactly the population the Students page counts.

GET /students/export/excel shares one filtered query with GET /students/ and
GET /students/search, so for the same filters the page total, the AJAX total
and the Excel row count are identical (pagination excluded).

Proves:
  * default "all except archived/deleted" — active/transferred/withdrawn/
    graduated exported, archived/deleted never; Arabic status labels;
  * a student whose stored academic_year_id is an older year is exported;
  * search / stage / grade / section / building / residential-area filters;
  * the whole filtered set, not the current 20-row page (page/per_page ignored);
  * School B students never reach School A's export, even with a tampered
    school_id; a crafted status=deleted exports nothing;
  * a building-restricted user and a teacher get the same scoped set in page
    and Excel; a tampered building_id cannot widen it;
  * the export link carries the current filters and never page;
  * a super admin with no school selected gets no cross-school file.
"""
import html
import re
import unittest
from datetime import date, datetime, timedelta
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app import create_app
from app.models import (db, AcademicYear, AuditLog, Employee, Grade, ResidentialArea,
                        Role, School, SchoolBuilding, Section, Student, User,
                        UserBuildingAccess)

PASSWORD = 'Test1234!'
OPTS = {'bypass_tenant_scope': True}
EXPORT = '/students/export/excel'
TOTAL_RE = re.compile(r'إجمالي (\d+) طالب')
EXPORT_LINK_RE = re.compile(r'<a href="([^"]+)" id="stu-export-link"')

ACTIVE = [f'A{i:02d}' for i in range(21)]                  # > one 20-row page
DEFAULT_VISIBLE = set(ACTIVE) | {'TR', 'WD', 'GR', 'OLD', 'NOS', 'S2A'}


class StudentsExportMatchesListTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.app = create_app('testing')
        cls.app.config['RATELIMIT_ENABLED'] = False
        with cls.app.app_context():
            cls.roles = {n: Role.query.filter_by(name=n).first().id
                         for n in ('school_admin', 'teacher', 'super_admin')}

    def setUp(self):
        self.sfx = uuid4().hex[:8]
        self.ids = {}
        with self.app.app_context():
            self._school('a')
            self._school('b')
            su = User(username=f'sxsu_{self.sfx}', email=f'sxsu_{self.sfx}@t.test',
                      full_name='su', role_id=self.roles['super_admin'], is_active=True)
            su.set_password(PASSWORD)
            db.session.add(su)
            db.session.commit()
            self.ids['super'] = su.username
            self.ids['super_id'] = su.id

    def _user(self, school, label, role):
        s = self.sfx
        u = User(username=f'sx{label}_{s}', email=f'sx{label}_{s}@t.test',
                 full_name=label, role_id=self.roles[role], school_id=school.id,
                 is_active=True)
        u.set_password(PASSWORD)
        db.session.add(u)
        db.session.flush()
        self.ids[label] = u.username
        self.ids[f'{label}_id'] = u.id
        return u

    def _school(self, key):
        s = self.sfx
        school = School(school_name=f'SXport {key} {s}', code=f'SX{key}{s}'[:20],
                        capacity=0, is_active=True, enable_buildings=(key == 'a'))
        db.session.add(school)
        db.session.flush()
        old = AcademicYear(school_id=school.id, name=f'Old{key}{s}', is_current=False,
                           start_date=date(2025, 8, 1), end_date=date(2026, 6, 30))
        cur = AcademicYear(school_id=school.id, name=f'Cur{key}{s}', is_current=True,
                           start_date=date(2026, 8, 1), end_date=date(2027, 6, 30))
        db.session.add_all([old, cur])
        db.session.flush()

        def section(name, stage, year):
            g = Grade(name=f'{name}{key}{s}', stage=stage, school_id=school.id,
                      academic_year_id=year.id)
            db.session.add(g)
            db.session.flush()
            sec = Section(name=f'{name}-1', grade_id=g.id, school_id=school.id,
                          academic_year_id=year.id)
            db.session.add(sec)
            db.session.flush()
            return g, sec

        g1, s1 = section('g1', 'ابتدائي', cur)
        g2, s2 = section('g2', 'متوسط', cur)
        _, s_old = section('gold', 'ابتدائي', old)
        self._user(school, f'adm{key}', 'school_admin')

        base = datetime(2026, 9, 1, 8, 0, 0)
        n = [0]

        def student(label, name, sec=None, status='active', year=cur, **extra):
            n[0] += 1
            db.session.add(Student(
                student_id=f'X{key.upper()}{s}-{label}', full_name=name,
                school_id=school.id, academic_year_id=year.id,
                section_id=sec.id if sec else None, status=status,
                created_at=base + timedelta(minutes=n[0]), **extra))

        if key == 'b':
            for label in ('A00', 'TR', 'OLD'):
                student(label, f'خالد طالب {label}', s1)
            db.session.flush()
            self.ids['school_b'] = school.id
            return

        b1 = SchoolBuilding(school_id=school.id, name=f'B1 {s}')
        b2 = SchoolBuilding(school_id=school.id, name=f'B2 {s}')
        area = ResidentialArea(school_id=school.id, name=f'Area {s}')
        db.session.add_all([b1, b2, area])
        db.session.flush()
        for i, label in enumerate(ACTIVE):
            student(label, f'طالب نشط {i:02d}', s1,
                    building_id=(b1 if i % 2 == 0 else b2).id,
                    residential_area_id=area.id if i < 3 else None)
        student('TR', 'خالد منقول', s1, 'transferred')
        student('WD', 'طالب مسحوب', s2, 'withdrawn')
        student('GR', 'طالب متخرج', s2, 'graduated')
        student('S2A', 'طالب الشعبة الثانية', s2)
        student('AR', 'طالب مؤرشف', s1, 'archived')
        student('DL', 'طالب محذوف', None, 'deleted')
        student('OLD', 'خالد سنة سابقة', s_old, year=old)      # older stored year
        student('NOS', 'طالب بلا شعبة', None)                   # NULL section

        # Building-restricted admin (B1 only) and a homeroom teacher of s2.
        restricted = self._user(school, 'rstr', 'school_admin')
        db.session.add(UserBuildingAccess(school_id=school.id, building_id=b1.id,
                                          user_id=restricted.id))
        teacher = self._user(school, 'tch', 'teacher')
        emp = Employee(school_id=school.id, employee_id=f'SXT{s}', full_name='T',
                       base_salary=0, status='active', user_id=teacher.id)
        db.session.add(emp)
        db.session.flush()
        s2.teacher_id = emp.id
        db.session.flush()
        self.ids.update(school_a=school.id, g1=g1.id, g2=g2.id, s1=s1.id, s2=s2.id,
                        b1=b1.id, b2=b2.id, area=area.id,
                        old_section=f'{s_old.grade.name} / {s_old.name}')

    def tearDown(self):
        with self.app.app_context():
            db.session.rollback()
            uids = [self.ids['super_id']]
            for key in ('a', 'b'):
                sid = self.ids[f'school_{key}']
                uids += [u.id for u in User.query.execution_options(**OPTS)
                         .filter_by(school_id=sid).all()]
            AuditLog.query.execution_options(**OPTS).filter(
                AuditLog.user_id.in_(uids)).delete(synchronize_session=False)
            for key in ('a', 'b'):
                sid = self.ids[f'school_{key}']
                for model in (UserBuildingAccess, Student, ResidentialArea, SchoolBuilding,
                              Section, Grade, Employee, User, AcademicYear):
                    model.query.execution_options(include_all_years=True, **OPTS).filter_by(
                        school_id=sid).delete(synchronize_session=False)
                School.query.filter_by(id=sid).delete(synchronize_session=False)
            User.query.execution_options(**OPTS).filter_by(
                id=self.ids['super_id']).delete(synchronize_session=False)
            db.session.commit()
            db.session.remove()

    # ── helpers ───────────────────────────────────────────────────────────────

    def _web(self, who='adma'):
        client = self.app.test_client()
        resp = client.post('/auth/login', data={'username': self.ids[who],
                                                'password': PASSWORD})
        self.assertEqual(resp.status_code, 302)
        return client

    def _labels(self, codes, key='A'):
        prefix = f'X{key}{self.sfx}-'
        self.assertTrue(all(c.startswith(prefix) for c in codes), codes)
        return {c[len(prefix):] for c in codes}

    def _sheet(self, client, **params):
        resp = client.get(EXPORT, query_string=params)
        self.assertEqual(resp.status_code, 200, resp.headers.get('Location'))
        self.assertEqual(resp.mimetype, 'application/vnd.openxmlformats-officedocument'
                                        '.spreadsheetml.sheet')
        ws = load_workbook(BytesIO(resp.data)).active
        return [r for r in ws.iter_rows(min_row=2, values_only=True)]

    def _assert_same(self, client, expected, key='A', **params):
        """Page total == AJAX total == Excel rows, and Excel rows == expected."""
        page = client.get('/students/', query_string=params)
        self.assertEqual(page.status_code, 200)
        page_total = int(TOTAL_RE.search(page.get_data(as_text=True)).group(1))
        ajax_total = client.get('/students/search', query_string=params).get_json()['total']
        rows = self._sheet(client, **params)
        codes = [r[1] for r in rows]
        self.assertEqual(len(codes), len(set(codes)))
        self.assertEqual(self._labels(codes, key), set(expected), params)
        self.assertEqual((page_total, ajax_total, len(rows)),
                         (len(expected),) * 3, params)
        return rows

    # ── 1, 2, 11: default population ──────────────────────────────────────────

    def test_default_page_total_equals_export_rows(self):
        rows = self._assert_same(self._web(), DEFAULT_VISIBLE)
        self.assertGreater(len(rows), 20)                    # not the 20-row page

    def test_default_includes_every_status_except_archived_and_deleted(self):
        rows = self._assert_same(self._web(), DEFAULT_VISIBLE)
        status_by = {r[1].rsplit('-', 1)[1]: r[8] for r in rows}
        self.assertEqual({k: status_by[k] for k in ('A00', 'TR', 'WD', 'GR')},
                         {'A00': 'فعال', 'TR': 'منقول', 'WD': 'مسحوب', 'GR': 'متخرج'})
        self.assertNotIn('AR', status_by)
        self.assertNotIn('DL', status_by)
        self.assertFalse({'active', 'transferred', 'withdrawn', 'graduated'}
                         & set(status_by.values()))

    def test_explicit_status_and_crafted_deleted(self):
        client = self._web()
        self._assert_same(client, ['AR'], status='archived')
        self._assert_same(client, ['TR'], status='transferred')
        self._assert_same(client, set(ACTIVE) | {'OLD', 'NOS', 'S2A'}, status='active')
        self.assertEqual(self._sheet(client, status='deleted'), [])

    def test_export_ignores_pagination_params(self):
        rows = self._sheet(self._web(), page=2, per_page=5, limit=3, offset=10)
        self.assertEqual(len(rows), len(DEFAULT_VISIBLE))

    # ── 3: older stored academic year ─────────────────────────────────────────

    def test_older_stored_year_student_is_exported_with_its_section(self):
        rows = self._assert_same(self._web(), DEFAULT_VISIBLE)
        old = next(r for r in rows if r[1].endswith('-OLD'))
        self.assertEqual(old[5], self.ids['old_section'])
        nos = next(r for r in rows if r[1].endswith('-NOS'))
        self.assertIn(nos[5], ('', None))

    # ── 4-7: page filters ─────────────────────────────────────────────────────

    def test_search_filter(self):
        client = self._web()
        self._assert_same(client, {'TR', 'OLD'}, q='خالد')
        self._assert_same(client, {'TR'}, q='خالد', status='transferred')
        self._assert_same(client, {'A05'}, q=f'XA{self.sfx}-A05')

    def test_stage_grade_section_filters(self):
        client = self._web()
        self._assert_same(client, {'WD', 'GR', 'S2A'}, stage='متوسط')
        self._assert_same(client, set(ACTIVE) | {'TR'}, grade_id=self.ids['g1'])
        self._assert_same(client, {'WD', 'GR', 'S2A'}, grade_id=self.ids['g2'],
                          section_id=self.ids['s2'])

    def test_building_filter(self):
        client = self._web()
        self._assert_same(client, set(ACTIVE[0::2]), building_id=self.ids['b1'])
        self._assert_same(client, set(ACTIVE[1::2]), building_id=self.ids['b2'])

    def test_residential_area_filter(self):
        self._assert_same(self._web(), {'A00', 'A01', 'A02'},
                          residential_area_id=self.ids['area'])

    # ── 8: tenant isolation ───────────────────────────────────────────────────

    def test_school_b_never_in_school_a_export(self):
        client = self._web('adma')
        for params in ({}, {'school_id': self.ids['school_b']}, {'q': 'خالد'}):
            with self.subTest(params=params):
                codes = [r[1] for r in self._sheet(client, **params)]
                self.assertTrue(codes)
                self.assertFalse([c for c in codes if c.startswith('XB')], codes)
        self._assert_same(self._web('admb'), {'A00', 'TR', 'OLD'}, key='B',
                          school_id=self.ids['school_a'])

    # ── 9, 10: restricted scopes ──────────────────────────────────────────────

    def test_building_restricted_user_same_population(self):
        client = self._web('rstr')
        b1_only = set(ACTIVE[0::2])
        self._assert_same(client, b1_only)
        # A building the user may not access cannot widen the set.
        self._assert_same(client, b1_only, building_id=self.ids['b2'])

    def test_teacher_scope_preserved(self):
        client = self._web('tch')
        self._assert_same(client, {'WD', 'GR', 'S2A'})
        self._assert_same(client, set(), grade_id=self.ids['g1'])

    # ── 12: export link ───────────────────────────────────────────────────────

    def test_export_link_carries_filters_not_page(self):
        client = self._web()
        body = client.get('/students/', query_string={
            'q': ' خالد ', 'status': 'transferred', 'page': 2,
            'building_id': self.ids['b1'], 'rfid': ''}).get_data(as_text=True)
        href = html.unescape(EXPORT_LINK_RE.search(body).group(1))
        self.assertTrue(href.startswith(EXPORT + '?'), href)
        from urllib.parse import parse_qs, urlsplit
        self.assertEqual(parse_qs(urlsplit(href).query),
                         {'q': ['خالد'], 'status': ['transferred'],
                          'building_id': [str(self.ids['b1'])]})
        # Unfiltered page → bare export URL (default = all except archived/deleted).
        body = client.get('/students/').get_data(as_text=True)
        self.assertEqual(html.unescape(EXPORT_LINK_RE.search(body).group(1)), EXPORT)
        self.assertNotIn("status='active'", body)
        # AJAX filtering rebuilds the link from getFilterParams() — no page.
        js = body[body.index('function exportHref()'):body.index('function doSearch(')]
        self.assertIn('getFilterParams()', js)
        self.assertNotIn('page', js)
        self.assertIn('exportLink.href = exportHref();', body)

    # ── super admin without a selected school ─────────────────────────────────

    def test_super_admin_without_school_gets_no_cross_school_file(self):
        resp = self._web('super').get(EXPORT)
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/students/', resp.headers['Location'])


if __name__ == '__main__':
    unittest.main()
